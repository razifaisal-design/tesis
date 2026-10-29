# -*- coding: utf-8 -*-
"""Recalc data/processed/Perhitungan_Hidrologi.xlsx (pustaka formulas) + verifikasi nilai
terhadap hasil_validasi.json, hasil_desain.json dan tabel neraca_bulanan di DB.
Jalankan setelah src/build_db.py -> src/analisis2.py -> src/analisis_desain.py ->
tools/buat_excel.py."""
import math
import json
import sqlite3
import sys
import formulas

from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "processed" / "hidrologi.db"
XLSX = ROOT / "data" / "processed" / "Perhitungan_Hidrologi.xlsx"
HV = json.load(open(ROOT / "data" / "processed" / "hasil_validasi.json"))
HD = json.load(open(ROOT / "data" / "processed" / "hasil_desain.json"))

print("[..] recalc dengan pustaka formulas ...")
xl = formulas.ExcelModel().loads(str(XLSX)).finish()
SOL = xl.calculate()
BOOK = XLSX.name.upper()
KEYMAP = {k.upper(): k for k in SOL if "!" in k}
NCHECK = NFAIL = 0
FAILS = []


def cval(sheet, ref):
    key = f"'[{BOOK}]{sheet.upper()}'!{ref}".upper()
    if key not in KEYMAP:
        raise KeyError(f"cell tidak ada di hasil recalc: {key}")
    v = SOL[KEYMAP[key]]
    try:
        return v.value[0, 0]
    except Exception:
        return v


def check(label, actual, expected, kind="rel", tol=1e-6):
    global NCHECK, NFAIL
    NCHECK += 1
    if kind == "str":
        ok = str(actual).strip() == str(expected)
    else:
        try:
            a, e = float(actual), float(expected)
        except (TypeError, ValueError):
            ok = False
        else:
            if kind == "rel":
                ok = abs(a - e) <= tol * max(abs(e), 1e-12)
            else:
                ok = abs(a - e) <= tol
    if not ok:
        NFAIL += 1
        FAILS.append(f"  GAGAL {label}: excel={actual!r} expected={expected!r}")
    return ok


# ---------- A. sheet lama (silang dgn hasil_validasi.json) ----------
check("Statistik!B4 Xbar", cval("Statistik", "B4"), HV["xbar"])
check("Statistik!B5 S", cval("Statistik", "B5"), HV["S"])
check("Statistik!B6 Cv", cval("Statistik", "B6"), HV["Cv"])
check("Statistik!B7 Cs", cval("Statistik", "B7"), HV["Cs"])
check("Statistik!B8 Ck", cval("Statistik", "B8"), HV["Ck"])
check("Statistik!B9 logXbar", cval("Statistik", "B9"), HV["logXbar"])
check("Statistik!B10 Slog", cval("Statistik", "B10"), HV["Slog"])
check("Statistik!B11 Cslog", cval("Statistik", "B11"), HV["Cslog"])
check("Frekuensi!G10 XT LP3 T=10", cval("Frekuensi", "G10"), HV["freq"]["10"]["LP3"], "abs", 0.001)
check("Chi_Kuadrat!B47 distribusi", cval("Chi_Kuadrat", "B47"), HV["chosen"], "str")
check("Debit!B16 R24", cval("Debit_Rencana", "B16"), HV["R24"], "abs", 0.0015)
check("Debit!B17 to", cval("Debit_Rencana", "B17"), HD["hidrolika"]["to"], "rel", 1e-4)
check("Debit!B18 v pipa", cval("Debit_Rencana", "B18"), HD["hidrolika"]["v"], "rel", 1e-4)
check("Debit!B19 td", cval("Debit_Rencana", "B19"), HD["hidrolika"]["td"], "rel", 1e-4)
check("Debit!B20 tc", cval("Debit_Rencana", "B20"), HD["hidrolika"]["tc"], "rel", 1e-4)
check("Debit!B21 I", cval("Debit_Rencana", "B21"), HD["hidrolika"]["I"], "rel", 1e-4)
check("Debit!B22 Q", cval("Debit_Rencana", "B22"), HD["hidrolika"]["Q"], "rel", 1e-4)
check("Debit!B23 Va", cval("Debit_Rencana", "B23"), HD["kolam"]["Va"], "rel", 1e-4)

# ---------- B. sheet Kawasan ----------
check("Kawasan!B19 Q kebutuhan (l/hari)", cval("Kawasan", "B19"), 76800.0, "abs", 1e-6)
check("Kawasan!B20 Q harian (m3/hari)", cval("Kawasan", "B20"), HD["Qharian"], "abs", 1e-9)

# ---------- C. sheet Hidrolika_Pipa ----------
H = HD["hidrolika"]
check("Hidrolika!B10 to Kirpich", cval("Hidrolika_Pipa", "B10"), H["to"])
check("Hidrolika!B14 R24", cval("Hidrolika_Pipa", "B14"), HV["R24"], "abs", 0.0015)
check("Hidrolika!B51 Q desain", cval("Hidrolika_Pipa", "B51"), H["Q"], "rel", 1e-4)
check("Hidrolika!B52 I desain", cval("Hidrolika_Pipa", "B52"), H["I"], "rel", 1e-4)
check("Hidrolika!B53 td", cval("Hidrolika_Pipa", "B53"), H["td"], "rel", 1e-4)
check("Hidrolika!B54 tc", cval("Hidrolika_Pipa", "B54"), H["tc"], "rel", 1e-4)
check("Hidrolika!B55 D teoritis (m)", cval("Hidrolika_Pipa", "B55"), H["Dteor"], "rel", 1e-4)
check("Hidrolika!B56 D komersial (mm)", cval("Hidrolika_Pipa", "B56"), H["Dkom"], "abs", 1e-9)
check("Hidrolika!B57 A pipa (m2)", cval("Hidrolika_Pipa", "B57"),
      math.pi * (H["Dkom"] / 1000) ** 2 / 4)
check("Hidrolika!B58 v pipa", cval("Hidrolika_Pipa", "B58"), H["v"], "rel", 1e-4)
check("Hidrolika!B59 kapasitas (Manning penuh)", cval("Hidrolika_Pipa", "B59"), H["kapasitas"], "rel", 1e-4)
check("Hidrolika!B60 cek kapasitas", cval("Hidrolika_Pipa", "B60"), "AMAN", "str")
check("Hidrolika!B61 D pelimpah", cval("Hidrolika_Pipa", "B61"), HD["kolam"]["D_overflow"], "abs", 1e-9)
U = HD["unit"]
check("Hidrolika!B66 Q unit", cval("Hidrolika_Pipa", "B66"), U["Q_unit"], "rel", 1e-4)
check("Hidrolika!B68 D komersial unit", cval("Hidrolika_Pipa", "B68"), U["D_kom"], "abs", 1e-9)
check("Hidrolika!B70 v unit", cval("Hidrolika_Pipa", "B70"), U["v"], "rel", 1e-4)
check("Hidrolika!B78 Q kapasitas talang", cval("Hidrolika_Pipa", "B78"), U["Q_talang"], "rel", 1e-6)
check("Hidrolika!B79 v talang", cval("Hidrolika_Pipa", "B79"), U["v_talang"], "rel", 1e-6)
check("Hidrolika!B80 V bak unit", cval("Hidrolika_Pipa", "B80"), U["V_bak"], "abs", 1e-6)

# ---------- C2. sheet Iklim_Harian + Evaporasi: rantai hitung Penman vs DB ----------
import csv as _csv
with open(ROOT / "data" / "processed" / "iklim_harian.csv") as _fh:
    IKL_CSV = list(_csv.DictReader(_fh))
for _idx in [0, len(IKL_CSV) // 2, len(IKL_CSV) - 1]:
    _rr = 4 + _idx
    check(f"Iklim_Harian!A{_rr} tanggal", cval("Iklim_Harian", f"A{_rr}"),
          IKL_CSV[_idx]["tanggal"], "str")
    check(f"Iklim_Harian!D{_rr} T", cval("Iklim_Harian", f"D{_rr}"),
          float(IKL_CSV[_idx]["t_c"]), "abs", 1e-9)
    check(f"Iklim_Harian!J{_rr} Rs", cval("Iklim_Harian", f"J{_rr}"),
          float(IKL_CSV[_idx]["rs_mj"]), "abs", 1e-9)

_con = sqlite3.connect(DB)
_con.row_factory = sqlite3.Row
EV = {r["bulan"]: r for r in _con.execute("SELECT * FROM evaporasi")}
_con.close()
EVA_MAP = [("C", "t_c"), ("D", "tmax_c"), ("E", "tmin_c"), ("F", "td_c"), ("G", "u10_ms"),
           ("H", "slp_hpa"), ("I", "rs_mj"), ("J", "es_kpa"), ("K", "ea_kpa"), ("L", "rh_pct"),
           ("M", "delta"), ("N", "gamma"), ("O", "u2_ms"), ("Q", "dr"), ("R", "deklinasi"),
           ("S", "omega_s"), ("T", "ra_mj"), ("U", "rso_mj"), ("V", "ratio"), ("W", "rns_mj"),
           ("X", "rnl_mj"), ("Y", "rn_mj"), ("Z", "lam"), ("AA", "ea_pen_mm"),
           ("AB", "e0_mm_hari"), ("AC", "evap_mm")]
for m in range(1, 13):
    rr = 13 + m
    check(f"Evaporasi!P{rr} J perwakilan", cval("Evaporasi", f"P{rr}"), EV[m]["j_hari"], "abs", 0)
    check(f"Evaporasi!AD{rr} n obs", cval("Evaporasi", f"AD{rr}"), EV[m]["n_obs"], "abs", 0)
    for col, dk in EVA_MAP:
        check(f"Evaporasi!{col}{rr} {dk}", cval("Evaporasi", f"{col}{rr}"), EV[m][dk], "rel", 1e-6)

# ---------- D. sheet Desain_Kolam ----------
K = HD["kolam"]
he = float(K["he"])
check("Kolam!B9 Ap iterasi konvergensi", cval("Desain_Kolam", "B9"), K["Ap"], "abs", 1e-6)
check("Kolam!B13 Va", cval("Desain_Kolam", "B13"), K["Va"], "rel", 1e-4)
check("Kolam!B16 Ap kriteria", cval("Desain_Kolam", "B16"), K["Ap_kriteria_Va"], "abs", 1e-6)
check("Kolam!B20 Vs", cval("Desain_Kolam", "B20"), K["Vs"], "abs", 0.05)
check("Kolam!B21 V desain", cval("Desain_Kolam", "B21"), K["V_desain"], "abs", 0.05)
check("Kolam!B22 l", cval("Desain_Kolam", "B22"), K["l"], "abs", 1e-9)
check("Kolam!B23 p", cval("Desain_Kolam", "B23"), K["p"], "abs", 1e-9)
check("Kolam!B24 Ap hasil dimensi live", cval("Desain_Kolam", "B24"), K["Ap"], "abs", 1e-6)
check("Kolam!B25 V kolam", cval("Desain_Kolam", "B25"), K["p"] * K["l"] * he, "abs", 1e-6)
check("Kolam!B26 h total", cval("Desain_Kolam", "B26"), K["htotal"], "abs", 1e-9)
check("Kolam!B41 selisih live-iterasi", cval("Desain_Kolam", "B41"), 0.0, "abs", 0.011)
check("Kolam!B42 status konvergensi",
      cval("Desain_Kolam", "B42"), "KONVERGEN (iterasi tuntas)", "str")
for i, th in enumerate(["0.5", "1", "1.5", "2", "3"]):
    check(f"Kolam!B{46+i} Va th={th}", cval("Desain_Kolam", f"B{46+i}"),
          K["Va_sensitivitas"][th], "rel", 1e-4)
FR = ["1.00", "0.70", "0.65", "0.60", "0.50"]
for i, fk in enumerate(FR):
    r = 54 + i
    check(f"Kolam!B{r} Vs f={fk}", cval("Desain_Kolam", f"B{r}"), HD["neraca"][fk]["vs"], "abs", 0.05)
    check(f"Kolam!C{r} reliability f={fk}", cval("Desain_Kolam", f"C{r}"),
          HD["neraca"][fk]["reliability"], "abs", 1e-6)
    check(f"Kolam!E{r} total net f={fk}", cval("Desain_Kolam", f"E{r}"),
          HD["neraca"][fk]["total_net"], "abs", 0.05)

# ---------- E. sheet Neraca_Air: seluruh deret vs tabel neraca_bulanan ----------
con = sqlite3.connect(DB)
dbn = {}
for y, b, f, sup, pond, dem, e, net, kum, defi, tp in con.execute(
        "SELECT tahun,bulan,fraksi,supply_m3,p_pond_m3,demand_m3,e_m3,net_m3,kumulatif_m3,defisit_m3,terpenuhi "
        "FROM neraca_bulanan"):
    dbn[(y, b, f)] = (sup, pond, dem, e, net, kum, defi, tp)
NB = len(list(con.execute("SELECT 1 FROM curah_hujan_bulanan")))
con.close()

from openpyxl import load_workbook
wbo = load_workbook(XLSX, data_only=False)
ws = wbo["Neraca_Air"]
blocks = []
for r in range(1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    if isinstance(a, str) and a.startswith("Fraksi f ="):
        blocks.append((float(a.split("=")[1].split()[0]), r))
assert len(blocks) == 5, f"blok fraksi ditemukan: {len(blocks)}"

COLS = [(6, 0, "supply"), (7, 1, "p_pond"), (8, 2, "demand"), (9, 3, "E"), (10, 4, "net"),
        (11, 5, "kumulatif"), (13, 6, "defisit"), (15, 7, "terpenuhi")]
NSERI = 0
for f, r0 in blocks:
    rd = r0 + 2
    for i in range(NB):
        rr = rd + i
        y = ws.cell(row=rr, column=1).value
        b = ws.cell(row=rr, column=2).value
        exp = dbn[(y, b, f)]
        for col, idx, nama in COLS:
            actual = cval("Neraca_Air", f"{chr(64 + col)}{rr}")
            NCHECK += 1
            NSERI += 1
            e = exp[idx]
            try:
                ok = abs(float(actual) - float(e)) <= (0 if nama == "terpenuhi" else 0.011)
            except (TypeError, ValueError):
                ok = False
            if not ok:
                NFAIL += 1
                if len(FAILS) < 40:
                    FAILS.append(f"  GAGAL Neraca {nama} ({y}-{b} f={f}): "
                                 f"excel={actual} db={e}")
    # ringkasan blok
    s = rd + NB
    fk = f"{f:.2f}"
    check(f"Neraca!E{s} Vs f={fk}", cval("Neraca_Air", f"E{s}"), HD["neraca"][fk]["vs"], "abs", 0.05)
    check(f"Neraca!E{s+1} reliability f={fk}", cval("Neraca_Air", f"E{s+1}"),
          HD["neraca"][fk]["reliability"], "abs", 1e-6)
    check(f"Neraca!E{s+3} total net f={fk}", cval("Neraca_Air", f"E{s+3}"),
          HD["neraca"][fk]["total_net"], "abs", 0.05)

# ---------- F. tidak boleh ada sel error formula ----------
NERR = 0
for key, v in SOL.items():
    if "!" not in key or key.startswith("'[") is False:
        continue
    try:
        val = v.value[0, 0]
    except Exception:
        continue
    if isinstance(val, str) and val.startswith("#") and len(val) <= 9:
        NERR += 1
        if NERR <= 20:
            FAILS.append(f"  ERROR formula di {key}: {val}")
NCHECK += 1
if NERR:
    NFAIL += 1

# ---------- laporan ----------
print(f"\n[cek] {NCHECK} pemeriksaan, {NFAIL} gagal "
      f"(termasuk {NSERI} nilai deret neraca vs DB)")
print(f"      deret Neraca_Air: 5 fraksi x {NB} bulan x 8 kolom (supply, hujan kolam, demand, E, net, kumulatif, defisit, terpenuhi)")
for line in FAILS[:40]:
    print(line)
if NFAIL:
    print("\n[HASIL] VERIFIKASI GAGAL")
    sys.exit(1)
print("\n[HASIL] VERIFIKASI LULUS - semua nilai Excel cocok dgn Python/DB")
print(f"  R24 = {cval('Hidrolika_Pipa', 'B14'):.3f} mm | Q = {cval('Hidrolika_Pipa', 'B51'):.4f} m3/dtk"
      f" | D = {cval('Hidrolika_Pipa', 'B56'):.0f} mm | v = {cval('Hidrolika_Pipa', 'B58'):.3f} m/dtk")
print(f"  Va = {cval('Desain_Kolam', 'B13'):.1f} | Vs = {cval('Desain_Kolam', 'B20'):.1f}"
      f" | V desain = {cval('Desain_Kolam', 'B21'):.1f} m3"
      f" -> {cval('Desain_Kolam', 'B23')} x {cval('Desain_Kolam', 'B22')} x {he} m")
print(f"  reliability: " + " | ".join(
    f"{fk}:{cval('Desain_Kolam', f'C{54+i}'):.1f}%" for i, fk in enumerate(FR)))
