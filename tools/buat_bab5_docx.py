#!/home/otherside/venv-project/bin/python
"""
Generator Naskah Resmi BAB V KESIMPULAN DAN SARAN (Word .docx & PDF)
Program Studi Magister Teknik Sipil - Konsentrasi Manajemen Rekayasa Sumber Daya Air
Universitas Tanjungpura, Pontianak.

Mengikuti standar format akademik:
- Margin: Kiri 4 cm, Atas 3 cm, Kanan 3 cm, Bawah 3 cm
- Font: Times New Roman 12 pt
- Line Spacing: 1,5 spasi
- Highlight kuning (WD_COLOR_INDEX.YELLOW) pada angka-angka kuantitatif kunci hasil rekonsiliasi
- Review marks [REVISI ANTIGRAVITY - HASIL INTEGRASI FASOS 814 m² & KELAYAKAN FINANSIAL]
"""

import os, subprocess
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DOC_BAB5_PATH = "/home/otherside/TESIS/BAB_V_KESIMPULAN_DAN_SARAN.docx"
PDF_BAB5_PATH = "/home/otherside/TESIS/BAB_V_KESIMPULAN_DAN_SARAN.pdf"
MODULAR_BAB5 = "/home/otherside/TESIS/Tesis_Lengkap/02_Bab_dan_Revisi_Parsial/BAB_V_KESIMPULAN_DAN_SARAN.docx"

print("=== MEMULAI GENERASI NASKAH BAB V KESIMPULAN DAN SARAN ===")

doc = docx.Document()

# Set Margins (Kiri 4 cm, Atas 3 cm, Kanan 3 cm, Bawah 3 cm)
section = doc.sections[0]
section.top_margin = Inches(3.0 / 2.54)
section.bottom_margin = Inches(3.0 / 2.54)
section.left_margin = Inches(4.0 / 2.54)
section.right_margin = Inches(3.0 / 2.54)
section.page_width = Inches(21.0 / 2.54)
section.page_height = Inches(29.7 / 2.54) # A4

def add_p(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5, highlight=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = bold
        run.italic = italic
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def add_run_hl(paragraph, text, bold=False, italic=False, highlight=True):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run

# ------------------------------------------
# JUDUL BAB V
# ------------------------------------------
p_title1 = add_p("BAB V", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=2)
p_title2 = add_p("KESIMPULAN DAN SARAN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

# ------------------------------------------
# 5.1 KESIMPULAN
# ------------------------------------------
add_p("5.1\tKesimpulan", bold=True, space_after=12)

p_intro = add_p(
    "Berdasarkan hasil analisis hidrologi, pemodelan hidraulika jaringan perpipaan, simulasi neraca air harian, perancangan geometri kolam retensi kriteria ganda di Lahan FASOS, serta evaluasi kelayakan tekno-ekonomi yang telah diuraikan pada Bab IV, dapat ditarik beberapa kesimpulan utama sebagai berikut:"
)

# Kesimpulan 1: Karakteristik Hujan & Potensi PAH
p_k1 = add_p("1.\t", bold=True)
p_k1.paragraph_format.left_indent = Inches(0.5)
p_k1.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_k1, "Potensi Pemanenan Air Hujan Kawasan: ", bold=True, highlight=False)
add_run_hl(p_k1, 
    "Kawasan Perumahan Subsidi Rajawali Residence di Desa Wonodadi 2, Kecamatan Sungai Raya, memiliki luas wilayah 2,88 hektar dengan 160 unit rumah Tipe 36. "
    "Total luas tangkapan atap bangunan mencapai ", highlight=False
)
add_run_hl(p_k1, "10.400 m²", bold=True, highlight=True)
add_run_hl(p_k1, " (1,04 ha atau 36,1% dari luas kawasan). Berdasarkan data resmi BPS Kabupaten Kubu Raya (Kubu Raya Dalam Angka), curah hujan tahunan rata-rata mencapai ", highlight=False)
add_run_hl(p_k1, "3.385,2 mm/tahun dengan 228 hari hujan/tahun", bold=True, highlight=True)
add_run_hl(p_k1, 
    " (didukung rata-rata Stasiun Meteorologi Supadio sebesar 3.080 mm/tahun). "
    "Dengan koefisien limpasan atap metal C = 0,80, potensi pasokan air hujan bruto tangkapan atap kawasan mencapai ", highlight=False
)
add_run_hl(p_k1, "28.163 m³/tahun", bold=True, highlight=True)
add_run_hl(p_k1, ", yang membuktikan bahwa sumber daya air hujan di kawasan ini sangat berlimpah dan potensial untuk dikonservasi.", highlight=False)

# Kesimpulan 2: Kebutuhan Domestik & Keandalan Pasokan (Reliability)
p_k2 = add_p("2.\t", bold=True)
p_k2.paragraph_format.left_indent = Inches(0.5)
p_k2.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_k2, "Kebutuhan Air Bersih dan Keandalan Pelayanan (Reliability): ", bold=True, highlight=False)
add_run_hl(p_k2, "Kebutuhan air bersih domestik non-potabel (cuci, sanitasi/flushing, dan siram) yang dilayani oleh sistem pemanenan air hujan ditetapkan sebesar ", highlight=False)
add_run_hl(p_k2, "60 liter/jiwa/hari (240 liter/hari per kepala keluarga)", bold=True, highlight=True)
add_run_hl(p_k2, ", sehingga kebutuhan harian total 160 unit rumah adalah ", highlight=False)
add_run_hl(p_k2, "38,4 m³/hari (14.016 m³/tahun)", bold=True, highlight=True)
add_run_hl(p_k2, ". Berdasarkan simulasi neraca air harian riil (daily water balance model) periode multi-tahun dengan kapasitas tampungan air baku efektif ", highlight=False)
add_run_hl(p_k2, "Vs = 516,6 m³ (~517 m³)", bold=True, highlight=True)
add_run_hl(p_k2, ", sistem mencapai tingkat keandalan pelayanan pasokan sebesar ", highlight=False)
add_run_hl(p_k2, "87,4%", bold=True, highlight=True)
add_run_hl(p_k2, " dengan volume air bersih yang berhasil dipasok mencapai ", highlight=False)
add_run_hl(p_k2, "12.082 m³/tahun", bold=True, highlight=True)
add_run_hl(p_k2, ". Kapasitas tampungan ini terbukti tangguh mengatasi periode defisit pada musim kemarau tipikal Kalimantan Barat hingga 14–20 hari tanpa hujan.", highlight=False)

# Kesimpulan 3: Efektivitas Atenuasi Banjir & Geometri Kolam FASOS 814 m²
p_k3 = add_p("3.\t", bold=True)
p_k3.paragraph_format.left_indent = Inches(0.5)
p_k3.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_k3, "Efektivitas Atenuasi Banjir dan Dimensi Realistis di Lahan FASOS: ", bold=True, highlight=False)
add_run_hl(p_k3, "Hujan rencana harian maksimum kala ulang 10 tahun (R10th) Stasiun Supadio terpilih menurut distribusi Log Pearson Tipe III adalah ", highlight=False)
add_run_hl(p_k3, "R24 = 153,5 mm", bold=True, highlight=True)
add_run_hl(p_k3, ". Untuk durasi badai standar drainase kawasan perumahan 1 jam (I = 53,22 mm/jam), debit puncak limpasan masuk atap total kawasan adalah ", highlight=False)
add_run_hl(p_k3, "Qin = 0,1231 m³/detik (123,1 L/s)", bold=True, highlight=True)
add_run_hl(p_k3, " dengan total volume limpasan masuk sebesar 442,8 m³. Dengan mendesain struktur pelimpah ambang (weir outflow) terkontrol sebesar ", highlight=False)
add_run_hl(p_k3, "Qout = 0,050 m³/detik (50 L/s)", bold=True, highlight=True)
add_run_hl(p_k3, " ke saluran drainase primer, kolam retensi kriteria ganda berhasil mereduksi debit banjir puncak kawasan sebesar ", highlight=False)
add_run_hl(p_k3, "59,3%", bold=True, highlight=True)
add_run_hl(p_k3, ", dengan volume tampungan sementara atenuasi banjir yang dibutuhkan sebesar ", highlight=False)
add_run_hl(p_k3, "Va = 350,0 m³", bold=True, highlight=True)
add_run_hl(p_k3, ".\n\nTotal kapasitas tampungan kolam kriteria ganda terpadu adalah ", highlight=False)
add_run_hl(p_k3, "Vdesain = Vs (517 m³) + Va (350 m³) + Vsedimen (55 m³) = 921,4 m³ (~876,8 m³ kapasitas muka air efektif)", bold=True, highlight=True)
add_run_hl(p_k3, ". Kolam retensi ini ", highlight=False)
add_run_hl(p_k3, "100% pas dan tepat dialokasikan pada kavling Lahan FASOS seluas 814,0 m² (22,0 m × 37,0 m)", bold=True, highlight=True)
add_run_hl(p_k3, 
    " sesuai Siteplan Resmi PBG 2022 (luas muka air tertinggi HWL 646 m², luas dasar kolam 450 m², kedalaman total 2,00 m, talud batu kali miring 1:1, dan cerucuk kayu dolken L = 4,0 m), "
    "sehingga memecahkan benturan spasial fatal pada draf awal dan sama sekali tidak menggusur kavling rumah subsidi warga di Row 6.",
    highlight=False
)

# Kesimpulan 4: Jaringan Perpipaan Kolektor & Distribusi
p_k4 = add_p("4.\t", bold=True)
p_k4.paragraph_format.left_indent = Inches(0.5)
p_k4.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_k4, "Desain Jaringan Perpipaan Kolektor dan Distribusi: ", bold=True, highlight=False)
add_run_hl(p_k4, "Jaringan pipa kolektor air hujan gravitasi dirancang pada bahu jalan sisi utara dengan pipa utama ", highlight=False)
add_run_hl(p_k4, "PVC SNI S-12.5 Ø 350 mm (14 inci)", bold=True, highlight=True)
add_run_hl(p_k4, " (kemiringan S = 0,5%, kapasitas penuh 148 L/s, kecepatan v = 1,54 m/s yang memenuhi syarat self-cleansing), pipa cabang blok ", highlight=False)
add_run_hl(p_k4, "PVC SNI Ø 200 mm (8 inci)", bold=True, highlight=True)
add_run_hl(p_k4, ", dan 16 titik manhole beton Ø 800 mm, mengoreksi rancangan awal pipa beton 36 inci yang over-engineered. Sedangkan jaringan distribusi air bersih bertekanan ditanam pada bahu jalan sisi selatan menggunakan pipa utama ", highlight=False)
add_run_hl(p_k4, "HDPE PE-100 PN-10 Ø 90 mm (3 inci)", bold=True, highlight=True)
add_run_hl(p_k4, " dan retikulasi blok ", highlight=False)
add_run_hl(p_k4, "HDPE Ø 63 mm (2 inci)", bold=True, highlight=True)
add_run_hl(p_k4, " yang dilayani oleh Rumah Pompa WTP (sand filter, activated carbon filter, desinfeksi sinar UV, dan ground reservoir 30 m³).", highlight=False)

# Kesimpulan 5: Kelayakan Finansial & Manfaat Sosial
p_k5 = add_p("5.\t", bold=True)
p_k5.paragraph_format.left_indent = Inches(0.5)
p_k5.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_k5, "Kelayakan Tekno-Ekonomi dan Manfaat Finansial Warga: ", bold=True, highlight=False)
add_run_hl(p_k5, "Total biaya investasi awal (CAPEX) sistem terpadu adalah ", highlight=False)
add_run_hl(p_k5, "Rp 953.757.000 (setara Rp 5.960.981 per unit rumah)", bold=True, highlight=True)
add_run_hl(p_k5, ", dan biaya operasional-pemeliharaan tahunan (OPEX) adalah ", highlight=False)
add_run_hl(p_k5, "Rp 28.608.000/tahun (setara Rp 14.900/bulan per KK)", bold=True, highlight=True)
add_run_hl(p_k5, ". Dengan penetapan iuran air bersih warga sebesar Rp 12.500/m³ (menghemat 64% dibanding tarif air tangki swasta Rp 35.000/m³), sistem menghasilkan penghematan biaya air bagi warga kawasan sebesar ", highlight=False)
add_run_hl(p_k5, "Rp 271.845.000 per tahun (hemat Rp 141.500/bulan per KK)", bold=True, highlight=True)
add_run_hl(p_k5, ". Pada umur teknis 20 tahun dan discount rate 6,5%, seluruh indikator kelayakan finansial terpenuhi dengan sangat baik:\n", highlight=False)

p_k5_sub = add_p(
    "a.\tNet Present Value (NPV) = + Rp 395.095.000  (NPV > 0 → Sangat Layak)\n"
    "b.\tBenefit-Cost Ratio (BCR) = 1,31  (BCR > 1,0 → Sangat Layak)\n"
    "c.\tPayback Period (PBP) = 7,79 tahun  (PBP jauh lebih pendek dari umur teknis 20 tahun)\n"
    "d.\tBreak-Even Point (BEP) tarif air impas = Rp 9.532 / m³  (< tarif applied Rp 12.500 / m³)",
    bold=True, space_after=18
)
p_k5_sub.paragraph_format.left_indent = Inches(0.75)
for r in p_k5_sub.runs:
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW

# ------------------------------------------
# 5.2 SARAN
# ------------------------------------------
add_p("5.2\tSaran", bold=True, space_after=12)

p_s_intro = add_p(
    "Berdasarkan hasil penelitian, temuan audit lapangan, dan kendala teknis yang dihadapi selama penyusunan tesis ini, beberapa saran konstruktif disampaikan untuk penerapan praktis maupun pengembangan riset lanjutan:"
)

# Saran 1: Aspek Konstruksi & Geoteknik
p_s1 = add_p("1.\t", bold=True)
p_s1.paragraph_format.left_indent = Inches(0.5)
p_s1.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_s1, "Aspek Pelaksanaan Konstruksi dan Geoteknik Tanah Lunak: ", bold=True, highlight=False)
add_run_hl(p_s1, 
    "Kondisi tanah kawasan studi yang bertopografi datar dengan formasi aluvial dan organosol/gambut memerlukan kehati-hatian tinggi pada saat pekerjaan galian kolam sedalam 2,0 meter. "
    "Pemancangan cerucuk kayu dolken Ø 10–12 cm sepanjang 4,0 meter dengan spasi 1,0 meter di bawah pondasi dinding talud batu kali wajib dilaksanakan secara seragam untuk mencegah keruntuhan geser lateral (slope failure) akibat gaya desak tanah jenuh air. "
    "Lantai dasar kolam harus diberi lapisan rabat beton kedap air t = 10–15 cm untuk menghindari rembesan air tanah gambut berkadar asam tinggi masuk ke dalam tampungan air baku.",
    highlight=False
)

# Saran 2: Operasional, Pemeliharaan & Kualitas Air
p_s2 = add_p("2.\t", bold=True)
p_s2.paragraph_format.left_indent = Inches(0.5)
p_s2.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_s2, "Operasional, Pemeliharaan Berkala, dan Jaminan Kualitas Air: ", bold=True, highlight=False)
add_run_hl(p_s2, 
    "Pengendapan sedimen dan sampah dedaunan dari atap harus ditangani secara preventif melalui pembersihan berkala kisi saringan (trash screen) pada manhole dan sediment forebay. "
    "Pengerukan endapan lumpur dead storage kolam retensi disarankan dilakukan minimal satu kali dalam setahun pada puncak musim kemarau (bulan Juli–Agustus) ketika volume muka air kolam berada pada titik terendah. "
    "Pemeriksaan kualitas air secara berkala (uji parameter pH, kekeruhan, besi Fe, mangan Mn, dan bakteri total coliform/E. coli) harus dilakukan di laboratorium lingkungan terakreditasi sebelum dan sesudah unit filtrasi WTP untuk memastikan air memenuhi baku mutu air higiene sanitasi Permenkes No. 2 Tahun 2023.",
    highlight=False
)

# Saran 3: Kelembagaan Pengelola Komunal & Partisipasi Warga
p_s3 = add_p("3.\t", bold=True)
p_s3.paragraph_format.left_indent = Inches(0.5)
p_s3.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_s3, "Kelembagaan Pengelola Komunal dan Tata Kelola Keuangan: ", bold=True, highlight=False)
add_run_hl(p_s3, 
    "Keberlanjutan sistem pemanenan air hujan terpadu sangat bergantung pada tata kelola kelembagaan di tingkat masyarakat. "
    "Disarankan agar pengembang menyerahkan pengelolaan prasarana air bersih dan kolam retensi kepada Unit Pengelola Air Bersih Komunal berbasis Rukun Warga (RW) atau Koperasi Warga Rajawali Residence. "
    "Penerapan meter air digital di setiap unit rumah dan penarikan iuran air berbasis volume pemakaian riil (Rp 12.500/m³) akan menciptakan transparansi, keadilan beban biaya, serta menjamin ketersediaan dana cadangan operasional (OPEX) dan pemeliharaan alat secara mandiri.",
    highlight=False
)

# Saran 4: Kebijakan Regulasi Pemerintah Daerah & Perizinan PBG
p_s4 = add_p("4.\t", bold=True)
p_s4.paragraph_format.left_indent = Inches(0.5)
p_s4.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_s4, "Rekomendasi Kebijakan bagi Pemerintah Daerah dan Pengembang: ", bold=True, highlight=False)
add_run_hl(p_s4, 
    "Dinas Pekerjaan Umum dan Penataan Ruang (PUPR) serta Dinas Perumahan Rakyat dan Kawasan Permukiman (Perkim) Kabupaten Kubu Raya disarankan menjadikan integrasi Kolam Retensi Kriteria Ganda pada Lahan FASOS sebagai pedoman teknis baku dalam proses validasi Persetujuan Bangunan Gedung (PBG) dan pengesahan rencana tapak (siteplan) perumahan baru. "
    "Pendekatan kriteria ganda ini terbukti menjadi solusi 'win-win' yang memecahkan dua masalah sekaligus: mengatasi defisit pasokan air bersih PDAM dan memitigasi risiko genangan banjir genangan di kawasan permukiman rawa gambut peri-urban.",
    highlight=False
)

# Saran 5: Pengembangan Riset Lanjutan
p_s5 = add_p("5.\t", bold=True)
p_s5.paragraph_format.left_indent = Inches(0.5)
p_s5.paragraph_format.first_line_indent = Inches(-0.5)
add_run_hl(p_s5, "Pengembangan Penelitian Akademis Lanjutan: ", bold=True, highlight=False)
add_run_hl(p_s5, 
    "Bagi peneliti selanjutnya, disarankan untuk: (a) Mengembangkan simulasi hidrodinamika numerik 2D/3D (misalnya menggunakan HEC-RAS 2D atau EPA SWMM) untuk memodelkan pola arus sedimentasi di dalam kolam secara lebih detail; "
    "(b) Mengintegrasikan sistem otomasi pemantauan ketinggian muka air dan katup pelimpah otomatis berbasis Internet of Things (IoT) dan sensor ultrasonik; serta "
    "(c) Melakukan studi Life Cycle Assessment (LCA) untuk mengukur jejak karbon (carbon footprint) penghematan energi listrik dari penggunaan air hujan dibandingkan sistem distribusi air perpipaan konvensional.",
    highlight=False
)

# Simpan berkas Word BAB V
doc.save(DOC_BAB5_PATH)
doc.save(MODULAR_BAB5)
print(f"Sukses menyimpan Naskah Word BAB V:\n - {DOC_BAB5_PATH} ({os.path.getsize(DOC_BAB5_PATH):,} bytes)\n - {MODULAR_BAB5}")

# Konversi ke PDF resmi via LibreOffice
print("Mengonversi naskah BAB V ke PDF resmi...")
subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", DOC_BAB5_PATH, "--outdir", "/home/otherside/TESIS/"], check=True)
print(f"Sukses ekspor PDF BAB V: {PDF_BAB5_PATH} ({os.path.getsize(PDF_BAB5_PATH):,} bytes)")

print("=== PENYUSUNAN BAB V KESIMPULAN DAN SARAN SELESAI SEMPURNA ===")
