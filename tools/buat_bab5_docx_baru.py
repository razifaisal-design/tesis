#!/usr/bin/env python3
"""Generator naskah BAB V KESIMPULAN DAN SARAN (baru) — docx.

Kesimpulan menjawab empat rumusan masalah Bab I; seluruh angka dari pipeline
tervalidasi (hasil_validasi.json, hasil_desain.json, DB). Draf lama
(BAB_V_KESIMPULAN_DAN_SARAN.docx) memakai angka proyek lama dan TIDAK dipakai.

Catatan: analisis kelayakan ekonomi (BEP) belum dihitung dari pipeline, sehingga
rumusan masalah ke-4 ditulis sebagai pekerjaan perancangan teknis lanjutan
(tidak mengarang angka NPV/BEP).

Output: manuscript/BAB_V_KESIMPULAN.docx
"""

import sys
from pathlib import Path

from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))

from naskah_util import (ROOT, DB, setup_doc, para, heading, bullet,
                         table, caption_table, P, set_run,
                         n, ni, rup, VAL, DES, KOL, HID, UNIT, demand_tahun)

OUT = ROOT / "manuscript" / "BAB_V_KESIMPULAN.docx"
doc = setup_doc()

Qharian = 76.8
dem65 = Qharian * 0.65 * 365

sup = DB.execute("SELECT SUM(supply_m3) s, SUM(p_pond_m3) p FROM neraca_bulanan "
                 "WHERE fraksi=0.65").fetchone()
NYEARS = len(VAL["years"])
supply_atap = sup["s"] / NYEARS
supply_kolam = sup["p"] / NYEARS


def butir(nomor, judul, teks):
    p = P(doc, after=8)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    set_run(p.add_run(f"{nomor}\t"), bold=True)
    set_run(p.add_run(judul + " "), bold=True)
    set_run(p.add_run(teks))
    return p


# ================================================================== JUDUL
p = P(doc, align=1, before=0, after=2)
set_run(p.add_run("BAB V"), bold=True)
p = P(doc, align=1, after=4)
set_run(p.add_run("KESIMPULAN DAN SARAN"), bold=True)

# ------------------------------------------------------------------ 5.1
heading(doc, "5.1", "Kesimpulan")
para(doc,
     "Berdasarkan hasil analisis hidrologi, perhitungan neraca air, dan penentuan "
     "dimensi kolam retensi serta jaringan pipa penyalur yang telah diuraikan pada "
     "Bab IV, dapat ditarik beberapa kesimpulan yang menjawab rumusan masalah "
     "penelitian sebagai berikut.")

butir("1.", "Perencanaan sistem pemanenan air hujan terpadu.",
      "Sistem dirancang sebagai satu kesatuan hidrolis yang menghubungkan intersepsi "
      "air hujan di atap rumah (talang atap) melalui jaringan pipa penyalur tertutup "
      "menuju bak penampung sementara per unit dan diteruskan ke kolam retensi "
      "komunal, untuk selanjutnya disalurkan ke unit rumah. Kawasan Perumahan Rajawali "
      "Residence (2,88 hektar, 160 unit rumah Tipe 36) memiliki luas atap total "
      "10.400 m\u00b2 (36,1% luas kawasan) dengan koefisien limpasan atap C = 0,80. "
      "Air hujan yang diintersepsi tidak dibuang sebagai limpasan, melainkan ditampung "
      "sebagai sumber air baku non-potabel, sekaligus menurunkan beban limpasan yang "
      "harus ditangani saluran drainase eksisting.")

butir("2.", "Kebutuhan air baku dan potensi tangkapan atap.",
      "Kebutuhan air domestik kawasan adalah " + n(Qharian, 1) + " m\u00b3/hari "
      "(" + ni(Qharian * 365) + " m\u00b3/tahun), dan kebutuhan desain pada fraksi "
      "non-potabel 65% adalah " + n(Qharian * 0.65, 2) + " m\u00b3/hari "
      "(" + ni(dem65) + " m\u00b3/tahun). Hujan rencana kala ulang 10 tahun "
      "(distribusi Log Pearson Tipe III) sebesar " + n(VAL["R24"], 3) + " mm, dengan "
      "hujan tahunan rata-rata 3.020,7 mm/tahun (Stasiun BWSK Sungai Ambawang, "
      "2012\u20132025). Potensi pasokan dari tangkapan atap mencapai "
      + ni(supply_atap) + " m\u00b3/tahun; dengan tambahan hujan langsung di "
      "permukaan kolam " + ni(supply_kolam) + " m\u00b3/tahun, total pemasukan "
      "menjadi " + ni(supply_atap + supply_kolam) + " m\u00b3/tahun, jauh melampaui "
      "kebutuhan desain. Secara volume tahunan, potensi tangkapan atap lebih dari "
      "cukup untuk memenuhi kebutuhan non-potabel kawasan; tingkat keandalan simulasi "
      "neraca bulanan 168 bulan pada fraksi 65% mencapai "
      + n(DES["neraca"]["0.65"]["reliability"], 1) + "%. Sebaliknya, untuk memenuhi "
      "100% kebutuhan domestik (termasuk air minum) keandalan hanya "
      + n(DES["neraca"]["1.00"]["reliability"], 1) + "%, sehingga sistem diposisikan "
      "secara tepat sebagai pemasok kebutuhan air non-potabel.")

butir("3.", "Dimensi talang, pipa, bak, kolam retensi, dan jaringan.",
      "Debit puncak limpasan atap pada kala ulang 10 tahun adalah " + rup(HID["Q"], 4)
      + " m\u00b3/detik (intensitas " + rup(HID["I"], 2) + " mm/jam pada tc = "
      + rup(HID["tc"] * 60, 2) + " menit), sehingga pipa penyalur utama ditetapkan "
      "DN" + ni(HID["Dkom"]) + " mm (kecepatan aliran penuh " + n(HID["v"], 3)
      + " m/detik; kapasitas " + n(HID["kapasitas"], 4) + " m\u00b3/detik, aman). "
      "Pipa penyalur per unit DN" + ni(UNIT["D_kom"]) + " mm (Q = "
      + n(UNIT["Q_unit"] * 1000, 2) + " liter/detik), talang atap berpenampang "
      "0,15 \u00d7 0,10 m (kapasitas " + n(UNIT["Q_talang"] * 1000, 2)
      + " liter/detik), dan bak penampung sementara per unit 0,5 m\u00b3. Volume "
      "tampungan ditentukan oleh kriteria ganda Vdesain = maks(Vs, Va): Vs (kebutuhan "
      "penyimpanan air baku, fraksi 65%) = " + rup(KOL["Vs"], 1) + " m\u00b3, "
      "sedangkan Va (atenuasi banjir, th = 1 jam) = " + rup(KOL["Va"], 1)
      + " m\u00b3, sehingga volume desain kolam = " + rup(KOL["V_desain"], 1)
      + " m\u00b3 (ditentukan fungsi penyediaan air baku). Dimensi kolam adalah "
      "panjang " + n(KOL["p"], 1) + " m \u00d7 lebar " + n(KOL["l"], 1)
      + " m \u00d7 kedalaman efektif " + n(KOL["he"], 1) + " m (ditambah freeboard "
      + n(KOL["hf"], 1) + " m; kedalaman total " + n(KOL["htotal"], 1)
      + " m), dengan luas permukaan Ap = " + n(KOL["Ap"], 2) + " m\u00b2.")

butir("4.", "Waktu pencapaian Break Even Point (BEP).",
      "Analisis kelayakan ekonomi (RAB, biaya operasional, pendapatan retribusi, dan "
      "BEP) berada pada lingkup perancangan teknis lanjutan dan belum dihitung dari "
      "basis data pipeline yang telah divalidasi. Angka-angka ekonomi pada draf "
      "sebelumnya bersumber dari basis data lama dan tidak dipakai. Perhitungan BEP "
      "perlu disusun ulang dengan mengacu pada dimensi dan volume hasil Bab IV "
      "(Vdesain = " + rup(KOL["V_desain"], 1) + " m\u00b3; pasokan tahunan "
      + ni(supply_atap + supply_kolam) + " m\u00b3) agar konsisten.")

para(doc, "Ringkasan angka utama hasil penelitian disajikan pada Tabel 5.1.", after=8)
caption_table(doc, "5.1", "Ringkasan Hasil Utama Penelitian")
table(doc, ["Besaran", "Nilai"],
      [["Hujan tahunan rata-rata (BWSK 2012\u20132025)", "3.020,7 mm/tahun"],
       ["Hujan rencana R24, T = 10 tahun (LP3)", n(VAL["R24"], 3) + " mm"],
       ["Intensitas pada tc", rup(HID["I"], 2) + " mm/jam (tc = "
        + rup(HID["tc"] * 60, 2) + " menit)"],
       ["Debit puncak limpasan atap", rup(HID["Q"], 4) + " m\u00b3/detik"],
       ["Diameter pipa penyalur utama", "DN" + ni(HID["Dkom"]) + " mm"],
       ["Kebutuhan air total", n(Qharian, 1) + " m\u00b3/hari"],
       ["Kebutuhan desain (fraksi 65%)", n(Qharian * 0.65, 2) + " m\u00b3/hari"],
       ["Pasokan atap tahunan", ni(supply_atap) + " m\u00b3/tahun"],
       ["Evaporasi kolam (Penman)", "1.847,2 mm/tahun"],
       ["Vs (fraksi 65%)", rup(KOL["Vs"], 1) + " m\u00b3"],
       ["Va (th = 1 jam)", rup(KOL["Va"], 1) + " m\u00b3"],
       ["Volume desain kolam = maks(Vs, Va)", rup(KOL["V_desain"], 1) + " m\u00b3"],
       ["Dimensi kolam", n(KOL["p"], 1) + " \u00d7 " + n(KOL["l"], 1) + " \u00d7 "
        + n(KOL["htotal"], 1) + " m"],
       ["Keandalan (fraksi 65%)", n(DES["neraca"]["0.65"]["reliability"], 1) + "%"]],
      widths=[7.0, 7.0])

# ------------------------------------------------------------------ 5.2
heading(doc, "5.2", "Saran")
para(doc,
     "Berdasarkan hasil penelitian dan kendala yang dihadapi, disampaikan beberapa "
     "saran untuk penerapan maupun pengembangan lanjutan.")
bullet(doc, "Verifikasi lapangan. Parameter yang masih berupa asumsi desain "
            "(panjang lereng atap Lo, kemiringan atap, kemiringan pipa, dan kedalaman "
            "efektif kolam) wajib diverifikasi melalui survei dan pengukuran "
            "waterpass, karena kemiringan pipa menentukan kecepatan aliran dan "
            "kapasitas penyaluran.")
bullet(doc, "Evaporasi dengan data lapangan. Hasil evaporasi Penman perlu divalidasi "
            "terhadap data panci evaporasi atau produk reanalisis lain, mengingat "
            "radiasi matahari bersumber dari produk satelit (NASA POWER).")
bullet(doc,
       "Sumber data hujan. Panjang data 14 tahun (2012\u20132025) memenuhi syarat "
       "minimum pedoman; penambahan panjang seri (>20 tahun) dianjurkan apabila data "
       "tersedia, agar analisis frekuensi lebih mantap.")
bullet(doc, "Alokasi lahan kolam. Luas permukaan kolam yang dibutuhkan (Ap = "
            + n(KOL["Ap"], 2) + " m\u00b2) perlu dipastikan ketersediaan lahannya pada "
            "siteplan kawasan, serta direkonsiliasi dengan rencana tata ruang dan "
            "lahan fasilitas umum/sosial setempat.")
bullet(doc, "Kajian ekonomi. Analisis kelayakan ekonomi (RAB, biaya operasional, "
            "retribusi, dan BEP) perlu disusun dari angka pipeline yang telah "
            "divalidasi, sejalan dengan rumusan masalah ke-4.")
bullet(doc, "Perancangan teknis lanjutan. Perancangan rinci kolam (talud, struktur "
            "pelimpah, pondasi pada tanah gambut), jaringan pipa distribusi "
            "bertekanan, dan pemilihan pompa (jalur kolam ke rumah warga) dilakukan "
            "pada tahap perancangan teknis berikutnya.")

doc.save(OUT)
print("Tersimpan:", OUT.relative_to(ROOT), f"({OUT.stat().st_size:,} bytes)")
