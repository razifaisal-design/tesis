"""Utilitas bersama untuk generator naskah tesis (docx) — Bab III, Bab IV,
lampiran, dan Bab V. Semua angka dibaca dari sumber kebenaran di data/processed.
"""

import json
import sqlite3
from pathlib import Path

import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
PROC = ROOT / "data" / "processed"
FIG = ROOT / "figures"
MAN = ROOT / "manuscript"

VAL = json.loads((PROC / "hasil_validasi.json").read_text(encoding="utf-8"))
DES = json.loads((PROC / "hasil_desain.json").read_text(encoding="utf-8"))
DB = sqlite3.connect(PROC / "hidrologi.db")
DB.row_factory = sqlite3.Row

KOL = DES["kolam"]
HID = DES["hidrolika"]
UNIT = DES["unit"]

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
         "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
BULAN3 = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun",
          "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]


# ------------------------------------------------------------------ format angka
def n(x, dec=2):
    if x is None:
        return "-"
    s = f"{x:,.{dec}f}"
    return s.replace(",", "\u2063").replace(".", ",").replace("\u2063", ".")


def ni(x):
    return n(x, 0)


def rup(x, dec=1):
    return n(x, dec)


def demand_tahun(fraksi):
    """Kebutuhan tahunan rata-rata (m3/th) = jumlah 168 bulan dibagi jumlah tahun."""
    s = DB.execute("SELECT SUM(demand_m3) s FROM neraca_bulanan WHERE fraksi=?",
                   (fraksi,)).fetchone()["s"]
    return s / len(VAL["years"])


def avg_bulan_kolom(tabel, kolom, bulan, fraksi=None):
    if fraksi is None:
        r = DB.execute(f"SELECT AVG({kolom}) a FROM {tabel} WHERE bulan=?", (bulan,)).fetchone()
    else:
        r = DB.execute(f"SELECT AVG({kolom}) a FROM {tabel} WHERE bulan=? AND fraksi=?",
                       (bulan, fraksi)).fetchone()
    return r["a"]


# ------------------------------------------------------------------ docx helpers
def setup_doc():
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.right_margin = Cm(3.0)
    sec.left_margin = Cm(4.0)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def set_run(r, bold=False, italic=False, size=12):
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return r


def P(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5,
      indent=None, first=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first is not None:
        pf.first_line_indent = Cm(first)
    return p


def para(doc, text="", bold=False, italic=False, **kw):
    p = P(doc, **kw)
    if text:
        set_run(p.add_run(text), bold=bold, italic=italic)
    return p


def heading(doc, num, text, before=14, after=8):
    p = P(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after)
    set_run(p.add_run(f"{num}\t{text}"), bold=True)
    return p


def subheading(doc, num, text, before=10, after=6):
    p = P(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after)
    set_run(p.add_run(f"{num}\t{text}"), bold=True)
    return p


def bullet(doc, text, level=0, marker="\u2022"):
    p = P(doc, after=4)
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.8)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    set_run(p.add_run(f"{marker}\t{text}"))
    return p


def caption_table(doc, num, judul):
    p = P(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=4, spacing=1.15)
    set_run(p.add_run(f"Tabel {num}"), bold=True)
    set_run(p.add_run(f" {judul}"))
    return p


def caption_gambar(doc, num, judul):
    p = P(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=12, spacing=1.15)
    set_run(p.add_run(f"Gambar {num}"), bold=True, size=11)
    set_run(p.add_run(f" {judul}"), size=11)
    return p


def gambar(doc, nama, num, judul, width_cm=14.0):
    path = FIG / nama
    if path.exists():
        p = P(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        para(doc, f"[Gambar {num} belum tersedia: {nama}]",
             align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    caption_gambar(doc, num, judul)


def shade(cell, hexfill="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def table(doc, headers, rows, widths=None, numerik=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j, h in enumerate(headers):
        hdr[j].text = ""
        p = hdr[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(str(h)), bold=True, size=size)
        shade(hdr[j])
    numerik = numerik or set()
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if j in numerik
                           else WD_ALIGN_PARAGRAPH.LEFT)
            set_run(p.add_run("" if v is None else str(v)), size=size)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
